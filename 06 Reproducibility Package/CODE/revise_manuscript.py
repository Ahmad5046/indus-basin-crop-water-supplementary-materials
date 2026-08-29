from __future__ import annotations

import shutil
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


ROOT = Path(r"E:\mdpi\indus_green_blue_water_vwf_20260709\Indus_Basin_Virtual_Water_Project")
REV = ROOT / "13_Q1_Comprehensive_Revision_20260718"
SOURCE = REV / "SOURCE_BACKUP" / "Indus_Basin_Manuscript_REVISED_equations-fixed_ORIGINAL.docx"
OUTPUT = REV / "Indus_Basin_Manuscript_Q1_REVISED.docx"
FIGURES = REV / "REVISED_FIGURES"
FONT = "Times New Roman"


def set_text(paragraph, text: str, style: str | None = None) -> None:
    paragraph.clear()
    paragraph.add_run(text)
    if style:
        paragraph.style = style


def find_prefix(doc: Document, prefix: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise KeyError(f"Paragraph not found: {prefix}")


def replace_prefix(doc: Document, prefix: str, text: str, style: str | None = None) -> None:
    set_text(find_prefix(doc, prefix), text, style)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_border(cell, **edges) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge, attrs in edges.items():
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in attrs.items():
            element.set(qn(f"w:{key}"), str(value))


def format_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    table.style = "Table Normal"
    set_repeat_table_header(table.rows[0])
    for ri, row in enumerate(table.rows):
        set_row_no_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(
                cell,
                top={"val": "nil"},
                left={"val": "nil"},
                bottom={"val": "nil"},
                right={"val": "nil"},
                insideH={"val": "nil"},
                insideV={"val": "nil"},
            )
            if ri == 0:
                set_cell_border(
                    cell,
                    top={"val": "single", "sz": "8", "color": "000000"},
                    bottom={"val": "single", "sz": "6", "color": "000000"},
                )
            if ri == len(table.rows) - 1:
                set_cell_border(cell, bottom={"val": "single", "sz": "8", "color": "000000"})
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1.5)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = FONT
                    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
                    run.font.size = Pt(8.5)
                    run.bold = ri == 0


def replace_embedded_image(docx_path: Path, member: str, image_path: Path) -> None:
    """Replace an embedded image while preserving the DOCX package structure."""
    with tempfile.TemporaryDirectory() as td:
        temp = Path(td) / docx_path.name
        with ZipFile(docx_path, "r") as src, ZipFile(temp, "w", ZIP_DEFLATED) as dst:
            for item in src.infolist():
                data = image_path.read_bytes() if item.filename == member else src.read(item.filename)
                dst.writestr(item, data)
        shutil.copy2(temp, docx_path)


doc = Document(SOURCE)

# Concise scope-accurate title and submission front matter.
set_text(doc.paragraphs[0], "Remote sensing-based green and blue crop water accounting in Pakistan's Indus Basin")
set_text(doc.paragraphs[1], "Running title: Crop water accounting in Pakistan's Indus Basin")
set_text(doc.paragraphs[2], "Author name(s) to be completed by the submitting author.")
set_text(doc.paragraphs[3], "Affiliation(s) to be completed by the submitting author.")
set_text(doc.paragraphs[4], "Corresponding-author details to be completed before submission.")
set_text(doc.paragraphs[5], "")
set_text(doc.paragraphs[6], "Submission note: complete the author, affiliation, e-mail, and ORCID fields in the journal system and manuscript before upload.", "Normal")

highlights = [
    "Spatial-block validation achieved 85.4% Rabi and 85.0% Kharif accuracy.",
    "Five crop classes occupied 17.80 million crop-hectares in the 2023 extent.",
    "Crop-season ET was 80.01 km3; wheat supplied 19.60 km3 residual blue ET.",
    "Cotton had the largest footprint and the lowest physical water productivity.",
    "Residual blue ET is an accounting estimate, not measured irrigation withdrawal.",
]
for paragraph, text in zip(doc.paragraphs[8:13], highlights):
    set_text(paragraph, text, "List Bullet")

abstract = (
    "Crop-specific consumptive water accounting is difficult in Pakistan's Indus Basin because cultivated area, "
    "evapotranspiration (ET), rainfall, and yield have different spatial and temporal supports. We integrated "
    "multisensor crop mapping with seasonal water accounting for cropland within the Pakistan portion of the Indus drainage "
    "basin in 2023; the boundary is not an official canal-command delineation. Separate Rabi and Kharif Random Forest models "
    "used Sentinel-2, Landsat 8/9, Sentinel-1, terrain, and ground-survey labels assigned to non-overlapping 0.05-degree "
    "training and validation blocks. Overall accuracies were 85.39% and 84.97%, respectively. Five target classes occupied "
    "17.80 million crop-hectares, with mapped-official area differences ranging from -34.66% for cotton to +61.95% for maize. "
    "FAO WaPOR v3 actual ET and interception was accumulated over crop calendars, while CHIRPS precipitation was converted to "
    "monthly effective rainfall using the USDA-SCS relationship. Green ET was operationally assigned as min(ET, effective "
    "rainfall); residual blue ET was max(ET - green ET, 0). This residual is not measured irrigation withdrawal or physical "
    "tracing of canal and groundwater sources. Crop-season ET totaled 80.01 km3. Wheat had the largest residual blue volume "
    "(19.60 km3; 74.06% dependency), whereas cotton had the largest total virtual water footprint (2076 m3 t-1) and lowest "
    "water productivity (0.482 kg m-3). Direct comparison with a MODIS-ERA5-Land chain produced substantially lower footprint "
    "endpoints, demonstrating strong structural sensitivity. Yield gaps, crop-area disagreement, a one-year crop extent, and "
    "the operational source partition limit causal interpretation. The workflow nevertheless provides an auditable screening "
    "basis for crop-season benchmarking and prioritizing field, canal-delivery, groundwater, salinity, and soil-water observations."
)
set_text(doc.paragraphs[14], abstract)
set_text(
    doc.paragraphs[15],
    "Keywords: actual evapotranspiration; crop mapping; green water; Indus Basin; residual blue water; virtual water footprint; water productivity",
)

# Sharpen the research gap and scope without introducing new claims.
replace_prefix(
    doc,
    "Previous work in the basin has established",
    "Previous work in the basin has established the value of satellite phenology, evapotranspiration mapping, and water accounting, but important integration gaps remain. Crop maps and consumptive-use estimates are often developed on different spatial supports or reported without showing how classification error affects water volume. Annual or generic seasonal totals can also obscure the contrasting calendars of Rabi wheat, monsoon rice, annual sugarcane, and dual-cycle maize. Finally, a spatially complete footprint map can conceal missing production denominators and can blur the distinction between classification performance, hydroclimatic product disagreement, and hydrological source attribution. Earlier Indus studies provide essential foundations for remote-sensing water accounting, yet an auditable chain from spatially validated crop labels through explicit calendars, seasonal water depth, mapped crop-hectares, and matched district yield remains uncommon (Cheema & Bastiaanssen, 2010; Karimi et al., 2013a; Usman et al., 2015; Siyal et al., 2023).",
)
replace_prefix(
    doc,
    "Against this background",
    "This study addresses those gaps for cropland within the Pakistan portion of the Indus drainage basin. The contribution is methodological integration rather than a new green-blue water theory: season-specific crop maps are evaluated with non-overlapping spatial holdout blocks; WaPOR ET and CHIRPS rainfall are accumulated over explicit crop calendars; each production-dependent indicator is traceable to mapped area and a positive matched district yield; missing yield remains non-computable; and disagreement between hydroclimatic product chains is reported separately from map accuracy. The objectives were to map wheat, rice, cotton, sugarcane, and maize; estimate an operational rainfall-attributed and residual ET partition; calculate crop, district, and provincial water-use, footprint, productivity, and dependency indicators; and identify relative priorities for follow-up. The outputs are screening evidence. They do not establish gross canal diversion, groundwater abstraction, legal command performance, or basin depletion.",
)

replace_prefix(
    doc,
    "Figure 1. Study area",
    "Figure 1. Pakistan portion of the Indus drainage-basin analysis boundary, provincial context, principal rivers, and district reporting units. The boundary combines HydroBASINS level-5 drainage geometry with FAO GAUL 2024 administrative units; spatial calculations use EPSG:32642. It is not an official canal-command boundary. Map lines delineate study areas and do not necessarily depict accepted national boundaries.",
    "Caption",
)
replace_prefix(
    doc,
    "The data sources were chosen",
    "The data sources were chosen to provide complementary information rather than interchangeable estimates (Table 1). Sentinel-2 supplied the principal optical phenology, Sentinel-1 added C-band backscatter, and Landsat 8/9 contributed a 30 m optical record (Drusch et al., 2012; Torres et al., 2012; Roy et al., 2014). Landsat Collection 2 processing followed U.S. Geological Survey documentation (USGS, 2021), while SRTM supplied elevation and slope (Farr et al., 2007). ESA WorldCover and Dynamic World provided complementary evidence for cropland extent (Zanaga et al., 2022; Brown et al., 2022), and the Pakistan ground-truth survey supplied crop labels (Ul Haq et al., 2024). WaPOR AETI was selected as the primary ET product on the basis of Food and Agriculture Organization of the United Nations documentation (FAO, 2020) and published evaluations (Blatchford et al., 2020; Seijger et al., 2023). CHIRPS served as the primary rainfall record because its station-satellite design is well documented and has been evaluated over Pakistan (Funk et al., 2015; Nawaz et al., 2021). Earth Engine provided the common processing environment for these large archives (Gorelick et al., 2017).",
)
replace_prefix(
    doc,
    "Crop calendars were defined",
    "Crop calendars were defined before water aggregation so that a single annual total was not assigned to crops with contrasting phenology (Table 2). Wheat was represented from 1 November of the preceding year to 30 April, rice from 1 June to 31 October, cotton from 1 April to 30 November, and sugarcane over the full calendar year. The maize window, 1 February to 30 November, intentionally spans the main spring and autumn cycles and therefore represents a combined seasonal class. These calendar choices structure the ET accounting rather than reproducing field-specific crop coefficients (Allen et al., 1998). They also affect area interpretation: mapped crop-hectares describe seasonal occupation and can count the same physical land in more than one season.",
)
replace_prefix(
    doc,
    "Official 2022-23 district crop yield",
    "Official 2022-23 district crop yield and production records were harmonized to GAUL district names (FAO, 2025b; Pakistan Bureau of Statistics, 2025). Mapped production equals mapped crop-season area multiplied by matched district yield and is therefore a hybrid remote-sensing/statistical estimate rather than the official production total. Five reported zero-yield rows were preserved, and records without positive matched yield remained non-computable. FAOSTAT national values were used only as an external consistency comparison (FAO, 2025a).",
)

# Reconcile reproducibility details with the archived execution artifacts.
replace_prefix(
    doc,
    "Separate Random Forest models were fitted",
    "Separate Random Forest models were fitted for Rabi and Kharif because class composition and phenological contrasts differ between seasons. The archived predictor stacks contain 42 variables for Rabi and 47 for Kharif, including period-specific Sentinel-2 indices; full-season Sentinel-2 and Landsat reflectance and indices; Sentinel-1 VV, VH, and VV-VH; and SRTM elevation and slope. Missing predictor pixels were filled with -0.25 to prevent cloud gaps from changing sample membership, although this sentinel value can influence tree splits. Both models used 300 trees, a minimum leaf population of 2, a bag fraction of 0.70, the default number of variables per split, no maximum-node constraint, and seed 20230709. This seed is recorded in the archived classification-accuracy JSON and is the authoritative value for the reported outputs. Random Forest was appropriate for the nonlinear mixed-sensor feature space (Breiman, 2001; Belgiu & Drăguţ, 2016), while evidence on sample size and composition informed the emphasis on independent validation (Schulthess et al., 2023).",
)
replace_prefix(
    doc,
    "To reduce spatial leakage",
    "To reduce spatial leakage, each 0.05-degree geographic block was assigned wholly to training or validation through a deterministic SHA-256 rule with seed 20230709; no block appeared in both subsets. The Rabi model used 5,535 training and 1,280 validation samples, while the Kharif model used 9,763 and 2,389. Confusion matrices provided overall accuracy, kappa, producer's accuracy, user's accuracy, F1, balanced accuracy, and macro-F1. Rabi balanced accuracy and macro-F1 were 0.846 and 0.855; the corresponding Kharif values were 0.855 and 0.852. Spatial separation makes the test more demanding than a point-wise random split, but it does not turn the class-balanced sample into an area-probability design. The accuracies therefore describe the validation sample, not the percentage of basin hectares mapped correctly, and Olofsson-type area adjustment is not defensible without design weights or a probability sample (Congalton, 1991; Olofsson et al., 2014).",
)

# Replace pseudo-Monte-Carlo interpolation with a direct structural envelope.
replace_prefix(
    doc,
    "Product-chain sensitivity was examined",
    "Product-chain structural sensitivity was examined by comparing two internally consistent endpoints: WaPOR AETI with CHIRPS rainfall as the primary chain and MODIS ET with ERA5-Land rainfall as the comparison chain. For each crop and indicator, both endpoint values, their absolute difference, and their percentage difference relative to the primary endpoint were reported directly. No random interpolation, probability distribution, or confidence interval was imposed. The comparison does not quantify classification, yield, crop-calendar, effective-rainfall, or aggregation uncertainty; it isolates only the consequence of substituting the complete ET-rainfall product chain.",
)
replace_prefix(
    doc,
    "X_A and X_B are the two product-chain endpoints",
    "X_A and X_B are the primary and comparison product-chain endpoints, respectively; ΔX_pct is the percentage difference relative to X_A. The difference is directional and has no probability or confidence-interval interpretation.",
)
replace_prefix(
    doc,
    "Quality control covered",
    "Quality control covered file existence and hashes, coordinate reference systems, raster dimensions and resolution, nodata behavior, units, negative values, water and footprint balance identities, district coverage, and cartographic layout. Recalculation from the authoritative crop table produced a maximum green-plus-blue water-balance residual of 3.83 m3 against approximately 80 billion m3, a maximum virtual-water residual of 2.22 x 10-7 m3 t-1, zero production-identity residual, and a maximum reciprocal water-productivity residual of 8.88 x 10-16 kg m-3. These differences are numerical precision, not hydrological imbalance. Maps were exported in PNG, TIFF, and PDF where available with Times New Roman typography. The revision audit preserves derived tables, software scripts, and SHA-256 lineage in a separate reproducibility package.",
)
replace_prefix(
    doc,
    "Note. All volumetric results use",
    "Note. All volumetric results use 1 mm over 1 ha = 10 m3. Equations 1-19 define the implemented indices, accounting relationships, diagnostic score, and direct endpoint comparison.",
)

# Improve result interpretation and add evidence-supported scenario diagnostics.
replace_prefix(
    doc,
    "Classification performance was similar across seasons",
    "Classification performance was similar across seasons. Rabi overall accuracy was 85.39% (kappa = 0.788; balanced accuracy = 0.846; macro-F1 = 0.855), based on 5,535 training and 1,280 spatially held-out validation samples. Wheat had high producer's accuracy (93.79%) but lower user's accuracy (79.78%), indicating greater commission than omission. Kharif overall accuracy was 84.97% (kappa = 0.811; balanced accuracy = 0.855; macro-F1 = 0.852), based on 9,763 training and 2,389 validation samples. Rice had the most balanced class performance, with producer's and user's accuracies of 92.46% and 91.40% (F1 = 0.919). Table 4 reports class-specific metrics, while Figure 4 shows the reconciled crop map used for water accounting. The results support regional pattern analysis, but the sample design does not support area-adjusted estimates or confidence intervals for mapped hectares.",
)
replace_prefix(
    doc,
    "Together, the five target classes represented",
    "Together, the five target classes represented 17.80 million crop-hectares. Wheat accounted for 7.93 million ha, followed by rice (4.16 million ha), maize (2.77 million ha), sugarcane (1.65 million ha), and cotton (1.29 million ha). These are season-specific occupied hectares, not 17.80 million unique physical hectares. Agreement with official harvested area varied sharply (Table 5): -6.45% for wheat, +40.00% for rice, -34.66% for cotton, +28.23% for sugarcane, and +61.95% for maize. Scaling mapped area to the official crop totals while holding mapped mean water depth and district-yield composition fixed changed crop water volume in direct proportion, for example from 26.47 to 28.29 km3 for wheat and from 19.63 to 12.12 km3 for maize. Because area and mapped production scale together in this diagnostic, crop-level VWF is unchanged. These are structural area scenarios rather than confidence limits; the discrepancies may reflect classification error, calendar and boundary mismatch, double cropping, perennial sugarcane, and the combined maize season.",
)
replace_prefix(
    doc,
    "Alternative product chains produced wide structural ranges",
    "The two product chains produced markedly different endpoints (Table 12; Figure 15). Total VWF declined from 1066.85 to 427.03 m3 t-1 for wheat (-59.97%), from 1548.60 to 413.74 for rice (-73.28%), from 2076.20 to 554.18 for cotton (-73.31%), from 125.63 to 38.61 for sugarcane (-69.27%), and from 1237.95 to 450.69 for maize (-63.59%). The corresponding comparison-chain water productivities were 2.342, 2.417, 1.804, 25.903, and 2.219 kg m-3, respectively, versus primary values of 0.937, 0.646, 0.482, 7.960, and 0.808. Cotton remained the largest-footprint crop and sugarcane the smallest under both chains, but the intermediate ordering changed. The result demonstrates strong product structural sensitivity; it is not an uncertainty interval and does not alter the designated primary estimates.",
)
replace_prefix(doc, "Table 12. Product-chain structural sensitivity", "Table 12. Direct two-endpoint structural sensitivity of total virtual water footprint (VWF, m3 t-1) and water productivity (WP, kg m-3).", "Caption")
replace_prefix(doc, "Note. Percentiles arise", "Note. The primary endpoint is WaPOR AETI with CHIRPS rainfall; the comparison endpoint is MODIS ET with ERA5-Land rainfall. Percentage differences are relative to the primary endpoint and are not confidence limits. VWF is reported in m3 t-1 and WP in kg m-3.")
replace_prefix(
    doc,
    "Figure 15. Structural sensitivity ranges",
    "Figure 15. Direct structural sensitivity of crop total virtual water footprint to the complete ET-rainfall product chain. Paired points show the primary WaPOR v3 AETI-CHIRPS endpoint and the MODIS ET-ERA5-Land comparison endpoint for the 2023 crop calendars. Connecting lines show product-chain displacement only; they are not confidence intervals.",
    "Caption",
)

# Discussion: distinguish map reliability, water accounting, management use, and evidence gaps.
replace_prefix(doc, "4.1 Interpreting the green-blue partition", "4.1 Reliability of the crop-mapping framework", "Heading 2")
replace_prefix(
    doc,
    "The green-blue partition becomes informative",
    "Spatial-block validation reduces the optimistic bias that can arise when neighboring survey points are divided randomly between model fitting and testing. The resulting overall accuracies near 85% are credible sample diagnostics, but they cannot be translated directly into correctly mapped hectares because the reference sample was class-balanced rather than selected with known area probabilities. Class-specific errors matter for the water account: wheat's lower user's than producer's accuracy indicates commission risk, while the larger crop-area disagreements for rice, cotton, sugarcane, and maize show that acceptable overall accuracy does not guarantee unbiased crop totals. The -0.25 gap-fill value, combined maize seasons, perennial sugarcane, and differences between mapped seasonal occupation and administrative harvested area remain plausible mechanisms. Accordingly, crop water volumes should be read together with Table 5 and the area-scaling scenarios in Supplementary Table S3.",
)
replace_prefix(doc, "4.2 Spatial concentration and hotspot meaning", "4.2 Interpretation of rainfall-attributed and residual consumptive use", "Heading 2")
replace_prefix(
    doc,
    "The provincial results show why",
    "The operational green-blue partition is informative when depth, area, and crop season are considered together. Wheat illustrates that interaction: despite the lowest crop-season ET depth, its extensive Rabi area and limited effective rainfall produced the largest residual blue volume and a blue share above 74%. Rice overlaps the monsoon and therefore received a larger rainfall-attributed share, but this does not imply low diversion, percolation, or pumping requirements. The rule assigns contemporaneous effective rainfall before residual ET; it does not resolve antecedent soil storage, capillary rise, canal seepage, return flow, irrigation timing, or the source of blue water. Residual blue ET is therefore a consistent screening metric for consumptive demand beyond modeled effective rainfall, not a physical measurement of irrigation withdrawal.",
)
replace_prefix(
    doc,
    "Hotspot status is deliberately relative",
    "Rainfall and ET uncertainty enter the partition nonlinearly because effective rainfall is estimated monthly and then capped by seasonal ET. CHIRPS evaluation over Pakistan supports regional use but also identifies limitations associated with gauge density and terrain (Nawaz et al., 2021); WaPOR and global ET assessments likewise show that retrieval structure and forcing data influence absolute values (Blatchford et al., 2020; Miralles et al., 2016; Pan et al., 2020). The large displacement between the two complete product chains is therefore scientifically consequential. Without independent flux, lysimeter, soil-water, and irrigation observations, neither endpoint can be treated as locally calibrated truth.",
)
replace_prefix(doc, "4.3 Comparison with previous research", "4.3 Crop-specific water use and productivity", "Heading 2")
replace_prefix(
    doc,
    "Comparison with earlier work is most defensible",
    "The crop results point to different mechanisms rather than a single basin-wide ranking. Wheat's importance is volumetric: its extensive area converts a moderate seasonal depth into 26.47 km3 of ET and 19.60 km3 of residual blue water. Cotton's concern lies in the production denominator: it combined the highest total and blue VWF with the lowest physical water productivity, but heat, pests, salinity, yield loss, and irrigation timing cannot be separated by the footprint alone. Rice had a lower residual-blue share because its calendar overlaps monsoon rainfall, yet its mapped area exceeded the official statistic by 40%, making the aggregate water total area-sensitive. Sugarcane had the greatest seasonal ET depth, while its apparently favorable kg m-3 value is dominated by fresh-stalk mass and cannot be compared directly with grain or cotton. Maize received the largest rainfall-attributed volume, but its broad February-November window blends spring and autumn cycles and likely contributes to the 62% mapped-official area difference.",
)
replace_prefix(
    doc,
    "The wide range reported in global crop-footprint studies",
    "These patterns are consistent with earlier Indus research showing strong spatial variation in ET and productivity and with water-accounting studies that distinguish consumptive use from withdrawals and recoverable return flows (Ahmad et al., 2009; Bastiaanssen et al., 2012; Karimi et al., 2013a; Simons et al., 2020; Usman et al., 2015). Direct numerical comparison is less defensible because published estimates differ in year, boundary, crop definition, calendar, ET algorithm, and production denominator. Irrigation strategy and hydrological representation can alter green-blue composition and footprint magnitude (Chukalla et al., 2015; Luan et al., 2018), while a low production denominator can raise m3 t-1 even when ET is not exceptional (Khan et al., 2021; Siyal et al., 2023). The present crop ordering therefore belongs to the adopted 2023 evidence chain rather than representing a universal agronomic ranking.",
)
replace_prefix(doc, "4.4 Irrigation and crop-planning implications", "4.4 Provincial and district contrasts and hotspot stability", "Heading 2")
replace_prefix(
    doc,
    "Crop-specific follow-up should match",
    "Provincial aggregation separates scale from source composition. Punjab dominates absolute crop-season ET because it contains most mapped crop-hectares, whereas Sindh has the highest residual-blue share. Balochistan and parts of Khyber Pakhtunkhwa rise in VWF rankings where consumed volume is divided by limited mapped production. The district results therefore identify two different follow-up problems: large blue volumes call for supply, timing, and consumed-fraction accounting, while high footprints call for crop-map and yield verification alongside agronomic diagnosis. The 75/25 diagnostic threshold selected 27 of 108 computable districts with a score of at least two; this count changed to 33, 22, and 11 under 70/30, 80/20, and 90/10 thresholds. Hotspot status is thus threshold- and denominator-sensitive, not evidence that an ecological or legal limit has been crossed.",
)
replace_prefix(
    doc,
    "Rice policy should likewise be spatially selective",
    "A minimum mapped-area screen further demonstrates denominator sensitivity. Restricting the analysis to districts with at least 25,000 mapped crop-hectares retained 76 of 108 computable districts and reduced the number scoring at least two under the 75/25 rule from 27 to 19. This does not establish a universally correct cutoff; it shows why small-area districts require explicit flags rather than automatic ranking. The relative priority index is most useful for selecting locations where field yields, salinity, groundwater conditions, canal service, and crop labels should be checked together.",
)
replace_prefix(doc, "4.5 Methodological strengths", "4.5 Comparison with previous research", "Heading 2")
replace_prefix(
    doc,
    "The workflow's principal strength is traceability",
    "Comparison with earlier work is strongest at the level of spatial patterns and accounting principles. Satellite analyses of the Indus likewise report pronounced variation in consumptive use and productivity (Ahmad et al., 2009; Bastiaanssen et al., 2012; Thenkabail, 2009). WA+ applications demonstrate the value of basin-consistent consumptive accounting (Karimi et al., 2013a, 2013b), while groundwater and consumed-fraction studies explain why ET cannot be equated with diversion or pumping (Cheema et al., 2014; Simons et al., 2020). Pakistan-focused footprint research similarly shows the interaction between water volume and production denominators (Khan et al., 2021; Siyal et al., 2023). Differences in analytical boundary, reference year, crop calendars, ET products, and yield statistics preclude direct equivalence among reported magnitudes.",
)
replace_prefix(
    doc,
    "Equally important is the treatment",
    "The distinctive value of the present workflow is traceability across evidence domains. Each crop result can be followed through a mapped class, validation design, crop calendar, ET product, rainfall product, effective-rainfall rule, and district-yield join. Earth Engine supports reproducible basin-scale processing, while continental crop-mapping experience illustrates both its scalability and its continuing dependence on reference data and validation design (Gorelick et al., 2017; Xiong et al., 2017). Missing yields remain visible instead of being filled silently, and mass-balance identities are tested. Crop-map accuracy is treated as part of water accounting rather than a separate cartographic exercise, while product structural sensitivity and production-denominator completeness are reported independently. This separation prevents unlike uncertainties from being collapsed into one confidence statement.",
)
replace_prefix(doc, "4.6 Structural sensitivity and interpretation", "4.6 Management relevance", "Heading 2")
replace_prefix(
    doc,
    "The structural-sensitivity ranges show",
    "The most defensible management use is a tiered monitoring-to-verification process. Seasonal crop maps and dekadal ET-rainfall diagnostics can identify crop zones or districts that depart from comparable seasonal behavior. Irrigation departments and crop-reporting services can then verify those signals with canal rotations and discharge, groundwater levels and salinity, crop-cutting observations, and field water balances; canal-performance evidence from Punjab shows why delivery context must accompany remote-sensing indicators (Sajid et al., 2022). Wheat follow-up should emphasize irrigation timing and conjunctive-use districts because small improvements across a large area may have substantial volumetric consequences. Rice and sugarcane zoning should be tested against monsoon reliability, canal service, aquifer condition, salinity, processing infrastructure, and farm income. Cotton priorities should combine water scheduling with yield-gap, heat, pest, and salinity diagnosis, while maize should be separated into spring and autumn cycles before season-specific recommendations are made.",
)
replace_prefix(
    doc,
    "Rainfall uncertainty enters the partition",
    "These diagnostics cannot support direct allocation or claims of water saving on their own. A reduction in diversion is not necessarily a reduction in consumptive use when seepage and return flows are recovered downstream, and residual blue ET cannot reveal whether the consumed water came from renewable surface supply or depleted aquifer storage. Command-level decisions require official service-area polygons, time-resolved deliveries, consumed-fraction and return-flow accounting, groundwater storage and salinity observations, and field verification. These requirements are consistent with broader Indus assessments that emphasize institutional, boundary, and allocation constraints (Laghari et al., 2012; Yang et al., 2013). Remote sensing narrows the search space; it does not replace the measurements needed for causal or regulatory decisions.",
)

replace_prefix(
    doc,
    "The hydroclimatic inputs introduce",
    "The hydroclimatic inputs introduce a separate layer of uncertainty. WaPOR AETI and CHIRPS rainfall are remotely sensed or model-assisted products, and the project contains no basin-wide flux-tower, lysimeter, soil-moisture, rain-gauge, or irrigation-delivery network for direct local validation. The MODIS-ERA5-Land endpoint reveals strong structural sensitivity but is not an uncertainty distribution. The short 2018-2022 fixed-extent comparison provides context rather than climatology because it holds the 2023 crop masks constant and cannot represent historical crop-area change.",
)
replace_prefix(
    doc,
    "Yield resolution limits",
    "Yield resolution limits the production side of the footprint. District-crop averages suppress within-district variation, so local footprint patterns reflect mapped water depth and crop class more strongly than field-scale yield. Of 580 district-crop combinations, 379 were computable and 201 lacked a positive matched yield or mapped crop; crop coverage ranged from 39.66% of districts for cotton to 93.10% for wheat. Five official zero-yield rows were retained rather than imputed. Islamabad and several small-area districts are particularly sensitive to small production denominators.",
)
replace_prefix(
    doc,
    "7. Alternative ET-rainfall chains",
    "7. Direct product-chain comparison showed that total VWF was 59.97-73.31% lower under the MODIS-ERA5-Land endpoint. Cotton and sugarcane retained the highest and lowest footprint positions, respectively, but intermediate ranks changed. The comparison demonstrates structural product sensitivity; it is not a confidence interval, and crop-map, calendar, effective-rainfall, and yield uncertainty remain additional components.",
)

# Honest availability and authorship statements.
replace_prefix(
    doc,
    "The derived datasets that support",
    "Derived datasets and reproducibility files have been assembled in a versioned local archive with a SHA-256 manifest but have not yet been published in a permanent repository. A public repository URL and DOI must be added before submission if the selected journal requires public deposition. Primary inputs remain with their original providers: Sentinel-1 and Sentinel-2, Landsat 8/9 Collection 2, SRTM, ESA WorldCover, Dynamic World, FAO WaPOR v3, CHIRPS, FAO GAUL 2024, WWF HydroBASINS, OpenStreetMap, the ADB/MNFSR/CRS Pakistan Agricultural Ground Truthing Survey (https://doi.org/10.5281/zenodo.11047562), and official Pakistan crop statistics. Provider terms continue to apply. No missing yield or field-validation values were imputed.",
)
replace_prefix(
    doc,
    "All code required to reproduce",
    "The Google Earth Engine and Python code, execution metadata, fixed random seeds, QC outputs, and derived tables have been assembled in the accompanying reproducibility package but are not yet publicly deposited. The package documents boundary and mask preparation, crop classification, spatial validation, seasonal ET and rainfall accounting, the operational green-blue partition, virtual-water calculations, direct product-chain comparison, cartography, tabulation, and quality control. A permanent repository URL and DOI should replace this statement before submission where required.",
)
replace_prefix(doc, "CRediT author statement.", "Author contributions will be finalized before submission.", "Normal")

# Insert an explicit AI-use declaration before References, as required by Elsevier.
references_heading = find_prefix(doc, "References")
ai_heading = OxmlElement("w:p")
ai_heading_pr = OxmlElement("w:pPr")
ai_style = OxmlElement("w:pStyle")
ai_style.set(qn("w:val"), "Heading1")
ai_heading_pr.append(ai_style)
ai_heading.append(ai_heading_pr)
ai_run = OxmlElement("w:r")
ai_text = OxmlElement("w:t")
ai_text.text = "14. Declaration of generative AI and AI-assisted technologies in the manuscript preparation process"
ai_run.append(ai_text)
ai_heading.append(ai_run)
ai_body = OxmlElement("w:p")
ai_body_run = OxmlElement("w:r")
ai_body_text = OxmlElement("w:t")
ai_body_text.text = (
    "During manuscript revision, the author used OpenAI Codex to assist with language editing, consistency checking, "
    "document automation, and reproducibility scripting. The author reviewed the resulting text and calculations and "
    "retains full responsibility for the content of the article."
)
ai_body_run.append(ai_body_text)
ai_body.append(ai_body_run)
references_heading._p.addprevious(ai_heading)
references_heading._p.addprevious(ai_body)

replace_prefix(
    doc,
    "Supplementary Figure S1 provides",
    "Supplementary files contain Figures S1-S2 and Tables S1-S10: full district rankings, validation diagnostics, crop-area scenarios, direct product-chain sensitivity, fixed-extent temporal context, hotspot stability, yield coverage, data provenance, scripts, QC reports, and the SHA-256 inventory.",
)

# Fix the style error in the note after Table 5.
replace_prefix(doc, "Note. Mapped production equals", find_prefix(doc, "Note. Mapped production equals").text, "Normal")

# Table 3: add Equation 19 to the symbol inventory.
table3 = doc.tables[2]
row = table3.add_row().cells
for cell, value in zip(row, ["ΔX, ΔX_pct", "Direct product-chain endpoint difference", "native unit; %", "Eq. 19"]):
    cell.text = value

# Table 12 is the thirteenth physical Word table (Tables 11a and 11b are separate).
table12 = doc.tables[12]
headers = [
    "Crop",
    "Primary VWF",
    "Comparison VWF",
    "Difference (%)",
    "Primary WP",
    "Comparison WP",
]
values = [
    ["Wheat", "1066.85", "427.03", "-59.97", "0.937", "2.342"],
    ["Rice", "1548.60", "413.74", "-73.28", "0.646", "2.417"],
    ["Cotton", "2076.20", "554.18", "-73.31", "0.482", "1.804"],
    ["Sugarcane", "125.63", "38.61", "-69.27", "7.960", "25.903"],
    ["Maize", "1237.95", "450.69", "-63.59", "0.808", "2.219"],
]
while len(table12.columns) < len(headers):
    table12.add_column(Inches(0.9))
while len(table12.rows) > 1:
    table12._tbl.remove(table12.rows[-1]._tr)
for j, text in enumerate(headers):
    table12.rows[0].cells[j].text = text
for vals in values:
    cells = table12.add_row().cells
    for j, text in enumerate(vals):
        cells[j].text = text

# Global journal-style formatting.
styles = doc.styles
for style_name in ["Normal", "Body Text", "List Bullet", "Caption"]:
    if style_name in styles:
        style = styles[style_name]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(12 if style_name != "Caption" else 10)
if "Table Normal" not in styles:
    styles.add_style("Table Normal", WD_STYLE_TYPE.TABLE)
for level, size in [(1, 14), (2, 12), (3, 12)]:
    style = styles[f"Heading {level}"]
    style.font.name = FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

for i, paragraph in enumerate(doc.paragraphs):
    text = paragraph.text.strip()
    if i == 0:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.size = Pt(16)
            run.bold = True
    elif paragraph.style.name == "Caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.keep_with_next = False
    elif paragraph.style.name.startswith("Heading"):
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.keep_with_next = True
    elif text:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = FONT
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        if paragraph.style.name == "Caption":
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0, 0, 0)
        elif not paragraph.style.name.startswith("Heading") and i != 0:
            run.font.size = Pt(12)

for table in doc.tables:
    format_table(table)

# Preserve source figure size and update Figure 15 without distortion.
new_fig = FIGURES / "Figure_15_Product_Chain_Endpoint_Sensitivity_1200dpi.png"
with Image.open(new_fig) as image:
    aspect = image.height / image.width
shape15 = doc.inline_shapes[14]
shape15.width = Inches(6.4)
shape15.height = int(shape15.width * aspect)

doc.core_properties.title = "Remote sensing-based green and blue crop water accounting in Pakistan's Indus Basin"
doc.core_properties.subject = "Q1-journal scientific revision for Agricultural Water Management"
doc.core_properties.comments = "Original preserved; revised from archived evidence without invented results."
doc.save(OUTPUT)
replace_embedded_image(OUTPUT, "word/media/image15.png", new_fig)

print(OUTPUT)
print(f"Abstract words: {len(abstract.split())}")
for text in highlights:
    print(len(text), text)
