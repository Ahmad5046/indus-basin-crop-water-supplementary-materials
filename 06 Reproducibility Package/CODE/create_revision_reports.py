from __future__ import annotations

import json
import platform
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"E:\mdpi\indus_green_blue_water_vwf_20260709\Indus_Basin_Virtual_Water_Project")
REV = ROOT / "13_Q1_Comprehensive_Revision_20260718"
TABLES = REV / "REVISED_TABLES"
FONT = "Times New Roman"


def base_doc(title: str, subtitle: str = "") -> Document:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
    for style_name in ["Normal", "Body Text", "List Bullet", "List Number", "Caption"]:
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = FONT
            style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
            style.font.size = Pt(10.5 if style_name != "Caption" else 9)
    for level, size in [(1, 14), (2, 12), (3, 11)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.name = FONT
    r.font.size = Pt(16)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(subtitle)
        r.italic = True
        r.font.name = FONT
        r.font.size = Pt(11)
    return doc


def add_p(doc: Document, text: str, *, bold_lead: str | None = None, bullet: bool = False) -> None:
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    for run in p.runs:
        run.font.name = FONT
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
        run.font.size = Pt(10.5)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def add_table(doc: Document, headers: list[str], rows: list[list[str]], font_size: float = 8.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_repeat_header(table.rows[0])
    for j, header in enumerate(headers):
        table.rows[0].cells[j].text = str(header)
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].text = str(value)
    for ri, row in enumerate(table.rows):
        set_no_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "E6E6E6" if ri == 0 else "FFFFFF")
            tc_pr.append(shd)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.name = FONT
                    r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
                    r.font.size = Pt(font_size)
                    r.bold = ri == 0
    doc.add_paragraph()


def fmt_value(value, name: str) -> str:
    if pd.isna(value):
        return "NA"
    if isinstance(value, str):
        return value
    if "count" in name or name.endswith("_code"):
        return f"{int(round(float(value))):,}"
    if "area_ha" in name or "production_t" in name:
        return f"{float(value):,.0f}"
    if "pct" in name or "accuracy" in name or "kappa" in name or "f1" in name:
        return f"{float(value):.3f}"
    return f"{float(value):.3f}"


def add_csv_table(doc: Document, title: str, path: Path, columns: list[str] | None = None, max_rows: int | None = None) -> None:
    doc.add_heading(title, level=2)
    frame = pd.read_csv(path)
    if columns:
        frame = frame[columns]
    if max_rows:
        frame = frame.head(max_rows)
    headers = [c.replace("_", " ").title() for c in frame.columns]
    rows = [[fmt_value(v, c) for c, v in zip(frame.columns, row)] for row in frame.itertuples(index=False, name=None)]
    add_table(doc, headers, rows, font_size=7.2 if len(headers) > 7 else 8.0)


# 1. Detailed change log.
change_log = base_doc("REVISION CHANGE LOG", "Comprehensive scientific reconstruction dated 18 July 2026")
change_log.add_heading("Purpose and scope", level=1)
add_p(change_log, "This log records substantive scientific, numerical, methodological, editorial, and presentation changes made to the preserved source manuscript. No primary result was changed without an archived source table or reproducible recalculation. Interpretive claims were narrowed where the evidence cannot support physical source attribution, command performance, groundwater abstraction, depletion, or regulatory conclusions.")
changes = [
    ["Title", "Long title implied complete IBIS command coverage and broad sustainability assessment.", "Replaced with a concise title referring to crop water accounting in Pakistan's Indus Basin.", "Matches the HydroBASINS drainage-basin cropland support.", "Source manuscript; boundary inventory", "No", "Resolved"],
    ["Abstract", "Exceeded AWM's 250-word limit and included pseudo-Monte-Carlo language.", "Rewritten to 249 words; direct product endpoint sensitivity and largest limitations retained.", "Complies with current AWM guidance and evidence.", "AWM Guide for Authors; project tables", "No", "Resolved"],
    ["Highlights", "Bullets exceeded journal length guidance.", "Five bullets reduced to 71-79 characters.", "AWM requires 3-5 bullets, each no more than 85 characters.", "AWM Guide for Authors", "No", "Resolved"],
    ["Study boundary", "Boundary could be interpreted as official canal commands.", "Defined as cropland within the Pakistan portion of the Indus drainage basin; map disclaimer added.", "Official command polygons were unavailable.", "HydroBASINS/GAUL processing files", "No", "Resolved by scope"],
    ["Crop classification", "Classifier seed stated as 20260709.", "Corrected to 20230709.", "Archived classification-accuracy JSON and block-split script record 20230709.", "classification_accuracy_2023.json; prepare_adb_ground_truth.py", "Yes: reproducibility metadata only", "Resolved"],
    ["Crop classification", "Predictor counts not explicit.", "Added 42 Rabi and 47 Kharif predictors and documented gap-fill behavior.", "Matches archived accuracy metadata.", "classification_accuracy_2023.json", "No", "Resolved"],
    ["Validation", "Only OA, kappa, and class metrics emphasized.", "Added balanced accuracy and macro-F1; clarified sample-based inference.", "Computed from archived confusion matrices.", "Table S1-S2", "Yes: added diagnostics", "Resolved"],
    ["Area inference", "Area adjustment could be inferred from validation.", "Explicitly rejected Olofsson adjustment because the class-balanced holdout is not an area-probability sample.", "Design weights and probability inclusion probabilities are absent.", "Ground-truth QC; Olofsson et al. (2014)", "No", "Resolved by limitation"],
    ["Crop area", "Mapped-official discrepancies were reported without volumetric scenarios.", "Added official-area-scaled water-volume scenarios while holding mapped depth/yield composition fixed.", "Shows first-order area sensitivity without treating it as confidence.", "Table S3", "Yes: scenario values", "Resolved"],
    ["ET", "Consumptive ET risked being read as delivered irrigation.", "Defined AETI as atmospheric consumptive loss; retained dekadal scale and prorating details.", "ET is not diversion or application.", "WaPOR metadata and scripts", "No", "Resolved"],
    ["Green-blue partition", "Physical source tracing could be inferred.", "Consistently renamed blue water as residual blue ET under a rainfall-first rule.", "Method omits storage, capillary rise, seepage, return flow, and source tracing.", "Equations 8-9; water scripts", "No", "Resolved by interpretation"],
    ["Yield", "Coverage limitations were incompletely quantified.", "Reported 379/580 computable district-crop records and crop-specific coverage.", "No imputation was introduced.", "Tables S9-S10", "Yes: added diagnostics", "Resolved"],
    ["Hotspots", "Single percentile rule could appear definitive.", "Added 70/30, 75/25, 80/20, and 90/10 sensitivity for all and >=25,000-ha districts.", "Quantifies threshold and denominator sensitivity.", "Tables S7-S8; Figure S2", "Yes: added diagnostics", "Resolved"],
    ["Structural sensitivity", "10,000 uniform draws between two endpoints implied probabilistic uncertainty.", "Removed interpolation; reported both endpoints, absolute/percentage differences, and rank changes.", "Only two complete product chains exist; no evidence-based input distributions are available.", "Tables S4-S5", "Yes: replaces artificial percentiles", "Resolved"],
    ["Equation 19", "Uniform interpolation equation no longer represented the method.", "Replaced with direct endpoint difference and percentage-difference equation in editable Word math.", "Matches revised sensitivity design.", "Table S5; manuscript method", "Yes: equation only", "Resolved"],
    ["Table 12/Figure 15", "Displayed constructed quantiles.", "Rebuilt as paired endpoint table and point-to-point structural comparison.", "Avoids false probabilistic interpretation.", "Table S4-S5; revised figure script", "Yes: presentation of supported endpoints", "Resolved"],
    ["Temporal context", "One-year results risked long-term interpretation.", "Added fixed-2023-extent 2018-2022 context in supplementary material only.", "Masks are fixed; values are not crop-area trends or a climatology.", "Table S6; trend manifest", "Yes: supplementary diagnostic", "Resolved by caveat"],
    ["Availability", "Manuscript claimed an open repository with a missing DOI.", "Replaced with an honest local-archive statement and pre-submission DOI requirement.", "No public deposition was verified.", "Project directory; no repository DOI", "No", "Unresolved until deposit"],
    ["AI declaration", "Elsevier disclosure absent.", "Added declaration describing Codex use for editing, consistency checks, automation, and scripting.", "Required transparency for actual tool use.", "AWM Guide for Authors", "No", "Resolved"],
    ["Formatting", "Captions were blue and sensitivity header crowded.", "Changed captions to black, shortened headers, preserved units in title/note, retained TNR.", "Improves journal-scale readability.", "Word COM visual QA", "No", "Resolved"],
]
add_table(change_log, ["Section", "Original problem", "Correction", "Scientific basis", "Supporting source", "Numerical effect", "Status"], changes, 7.2)
change_log.add_heading("Numerical integrity statement", level=1)
add_p(change_log, "The primary 2023 crop areas, ET depths, green and residual-blue depths, water volumes, mapped production, virtual water footprints, water productivity, province summaries, and leading district values were not changed. New numbers are diagnostics derived from archived outputs: balanced accuracy, macro-F1, area-scaling scenarios, direct product-chain endpoint differences, fixed-extent temporal context, hotspot threshold sensitivity, and yield coverage.")
change_log.save(REV / "REVISION_CHANGE_LOG.docx")


# 2. Scientific audit.
audit = base_doc("SCIENTIFIC METHODS AUDIT", "Remote sensing-based green and blue crop water accounting")
audit.add_heading("Executive assessment", level=1)
add_p(audit, "The implemented arithmetic is internally consistent, and the revised manuscript now describes the executed workflow with substantially better fidelity. The strongest evidence consists of spatial-block crop validation, explicit calendars, product-specific ET/rainfall accounting, traceable district-yield joins, and mass-balance QC. The principal weakness is not a calculation error but evidential support: residual blue ET is an operational accounting residual, ET/rainfall products lack local field validation, crop areas disagree substantially with official harvested areas, and yield is incomplete and district-aggregated.")
audit.add_heading("Issue inventory", level=1)
issues = [
    ["Critical", "Probabilistic uncertainty claim from two endpoint interpolation", "Replaced with direct endpoint structural sensitivity", "Resolved"],
    ["Critical", "Residual blue ET could be interpreted as measured irrigation withdrawal", "Explicitly excluded diversion, pumping, depletion, and physical source tracing", "Resolved by scope"],
    ["Major", "Classifier seed conflict", "Archived output provenance establishes 20230709; later 20260709 code copies flagged", "Resolved for manuscript; code lineage documented"],
    ["Major", "Large mapped-official crop-area disagreements", "Retained, discussed, and propagated through first-order area scenarios", "Partly quantified; mechanism unresolved"],
    ["Major", "Validation is not area-probability based", "No area-adjusted estimate or CI claimed", "Unresolved without new sample design"],
    ["Major", "No local field validation for ET/rainfall/partition", "Product evaluation literature and direct endpoint sensitivity used", "Unresolved without observations"],
    ["Major", "Missing/non-positive district yields", "No imputation; 379/580 combinations computable", "Transparent but unresolved"],
    ["Moderate", "Combined spring/autumn maize and annual sugarcane", "Calendar limitations tied to area mismatch and interpretation", "Unresolved without separate mapping"],
    ["Moderate", "Hotspot thresholds and small denominators", "Threshold and >=25,000-ha screens added", "Quantified sensitivity"],
    ["Moderate", "One-year crop extent", "Fixed-extent multi-year context labelled non-climatological", "Unresolved for interannual crop dynamics"],
    ["Minor", "Caption color, table density, equation rendering", "Corrected and visually verified in Word/PDF", "Resolved"],
]
add_table(audit, ["Severity", "Issue", "Treatment", "Disposition"], issues, 8.0)
audit.add_heading("Crop classification and validation", level=1)
add_p(audit, "Separate Rabi and Kharif Random Forest classifiers use mixed optical, radar, and terrain predictors. The reported 0.05-degree block split reduces direct neighborhood leakage. The Rabi and Kharif confusion matrices yield OA values of 0.854 and 0.850, kappa values of 0.788 and 0.811, balanced accuracies of 0.846 and 0.855, and macro-F1 values of 0.855 and 0.852. These metrics describe the held-out sample. Because the reference design is class-balanced and lacks area inclusion probabilities, it does not support unbiased area estimates or Olofsson-type area adjustment.")
audit.add_heading("ET and rainfall processing", level=1)
add_p(audit, "WaPOR v3 L1_AETI_D is treated as a dekadal mean daily rate, multiplied by the catalog scale factor and intersecting calendar days. Partial dekads are prorated. CHIRPS Daily rainfall is aggregated monthly before applying the nonlinear USDA-SCS effective-rainfall relationship and is then summed over the crop calendar. The calculations are dimensionally coherent, but local validation and a full root-zone water balance are absent.")
audit.add_heading("Green-blue partition", level=1)
add_p(audit, "The identities ET_green = min(ET, P_eff) and ET_blue = max(ET - ET_green, 0) are mathematically correct for the adopted rainfall-first convention. They do not physically identify the origin of consumed water. Antecedent storage, pre-season rainfall, capillary rise, groundwater uptake, canal seepage, return flow, rooting depth, runoff, deep percolation, and rainfall-irrigation interactions are outside the model. The manuscript now uses 'rainfall-attributed ET' and 'residual blue ET' where physical source language would overstate the evidence.")
audit.add_heading("Footprint, productivity, and unit audit", level=1)
for text in [
    "1 mm over 1 ha equals 10 m3; the conversion is correct.",
    "Mapped production equals mapped crop-hectares multiplied by matched district yield in t ha-1.",
    "Crop/province/district VWF uses summed water volume divided by summed production; local ratios are not averaged.",
    "WP uses production mass in kg divided by consumptive water volume in m3 and is the reciprocal of total VWF after the tonne-to-kilogram conversion.",
    "Sugarcane fresh-stalk mass is not directly comparable with grain or cotton product mass.",
    "VWF and WP are not irrigation efficiency, economic productivity, scarcity, depletion, or sustainability metrics.",
]:
    add_p(audit, text, bullet=True)
audit.add_heading("Numerical reconciliation", level=1)
qc = json.loads((TABLES / "scientific_audit_qc.json").read_text(encoding="utf-8"))
qc_rows = [[k.replace("_", " "), str(v)] for k, v in qc.items()]
add_table(audit, ["Check", "Result"], qc_rows, 8.0)
audit.add_heading("Audit verdict", level=1)
add_p(audit, "The revised analysis is defensible as a screening-level remote-sensing water account. It is not sufficient for claims about irrigation withdrawal, canal performance, groundwater abstraction, aquifer sustainability, environmental-flow compliance, or actual water saving. Those boundaries are now explicit in the title, abstract, methods, discussion, limitations, and conclusion.")
audit.save(REV / "SCIENTIFIC_METHODS_AUDIT.docx")


# 3. Journal readiness report.
ready = base_doc("JOURNAL READINESS REPORT", "Primary target: Agricultural Water Management; secondary target: Remote Sensing Applications: Society and Environment")
ready.add_heading("Journal-fit decision", level=1)
add_p(ready, "Agricultural Water Management score: 6.5/10. The article has a strong applied water-accounting question, a substantial spatial analysis, transparent limitations, and direct management relevance. Its main barriers are the lack of local ET/soil-water/irrigation validation, the large crop-area discrepancies, a one-year crop extent, and the operational rather than physical source partition.")
add_p(ready, "Remote Sensing Applications: Society and Environment score: 8.0/10. The multisensor crop mapping, Earth Engine implementation, spatial validation, crop-calendar accounting, and decision-screening emphasis align more directly with this journal. The same limitations remain, but they are less likely to be decisive if framed as an application and reproducibility study.")
add_p(ready, "Recommended submission route: Remote Sensing Applications: Society and Environment in the present evidence state. Agricultural Water Management becomes more realistic after independent field or command-level validation, probability-based area assessment, separate maize-season mapping, and public deposition of the code/data archive.")
ready.add_heading("Agricultural Water Management requirements", level=1)
checks = [
    ["Single-column editable Word manuscript", "Met", "DOCX is editable and Word 16.0 generated the PDF."],
    ["Abstract <=250 words", "Met", "249 words."],
    ["Highlights: 3-5 bullets, <=85 characters", "Met", "Five bullets, 71-79 characters."],
    ["Keywords: 1-7", "Met", "Seven keywords."],
    ["Editable numbered equations", "Met", "Nineteen Word OMath equations, cited sequentially."],
    ["Tables and figures cited sequentially", "Met", "Fifteen figures and twelve numbered table captions, with 11a/11b distinguished."],
    ["Graphical abstract", "Met as separate asset", "AWM-ratio PNG/TIFF/PDF created; upload separately."],
    ["Data availability", "Conditional", "Honest local-archive statement; public DOI remains unresolved."],
    ["Author/affiliation/corresponding author", "Not met", "Submitting author must supply identities and metadata."],
    ["Declaration of interests/funding", "Met subject to author confirmation", "Neutral statements retained."],
    ["AI-use declaration", "Met", "Actual use disclosed before References."],
    ["High-resolution artwork", "Met for supplied assets", "Raster/vector files retained; scientific detail remains limited by native source resolution."],
]
add_table(ready, ["Requirement", "Status", "Evidence/action"], checks, 8.0)
ready.add_heading("Desk-stage blockers", level=1)
for text in [
    "Complete author names, affiliations, corresponding-author e-mail, and ORCID where applicable.",
    "Deposit the approved reproducibility archive and replace the conditional data/code statements with a verified DOI/URL.",
    "Confirm whether the target journal wants the graphical abstract and highlights as separate upload files only.",
    "Confirm author contributions, funding, conflict-of-interest wording, and acknowledgements.",
]:
    add_p(ready, text, bullet=True)
ready.add_heading("Current official guidance consulted", level=1)
add_p(ready, "Agricultural Water Management Guide for Authors: https://www.sciencedirect.com/journal/agricultural-water-management/publish/guide-for-authors")
add_p(ready, "Remote Sensing Applications: Society and Environment Guide for Authors: https://www.sciencedirect.com/journal/remote-sensing-applications-society-and-environment/publish/guide-for-authors")
ready.save(REV / "JOURNAL_READINESS_REPORT.docx")


# 4. Unresolved gaps.
gaps = base_doc("UNRESOLVED EVIDENCE GAPS", "Inputs or evidence needed beyond the completed project")
gap_rows = [
    ["Author metadata", "No verified names, affiliations, e-mail, or ORCID in project", "Desk-stage submission fields cannot be completed", "Submitting author supplies and confirms metadata", "Pre-submission"],
    ["Repository DOI", "Archive exists locally but no verified public DOI", "Data/code availability cannot claim public access", "Deposit approved archive in Zenodo or institutional repository", "Pre-submission where required"],
    ["Area-adjusted crop estimates", "Validation is not an area-probability sample", "No unbiased mapped-area CI", "Stratified probability sample with inclusion probabilities/design weights", "High"],
    ["Province-specific validation", "Archived validation predictions lack spatially attributed records", "Cannot quantify regional classifier reliability", "Export point-level validation predictions with coordinates/blocks", "High"],
    ["ET validation", "No flux-tower, lysimeter, or field water-balance network", "Absolute ET endpoint accuracy is unknown locally", "Representative field ET observations across agroecological zones", "High"],
    ["Rainfall validation", "No project rain-gauge archive linked to crop seasons", "Local CHIRPS bias cannot be corrected", "Quality-controlled gauges and elevation-stratified validation", "High"],
    ["Physical source partition", "No root-zone storage, irrigation timing, capillary rise, or seepage accounting", "Residual blue ET cannot identify canal versus groundwater", "Soil-water balance or assimilated hydrological model with observations", "High"],
    ["Canal-command analysis", "Official command polygons and time-resolved deliveries unavailable", "No command performance or allocation inference", "Official service areas, rotations, gate/discharge records, return-flow data", "High"],
    ["Groundwater sustainability", "No abstraction, storage change, water level, recharge, or salinity series", "No pumping or depletion attribution", "Well hydrographs, abstraction/energy records, storage and salinity monitoring", "High"],
    ["Crop-calendar regionalization", "One basin-wide calendar per crop; maize combines two cycles", "ET/rainfall aggregation may be seasonally mismatched", "Province/agroecological crop calendars and separate spring/autumn maize maps", "Medium"],
    ["Yield resolution", "District means and 201 non-computable district-crop rows", "Local VWF and hotspot ranks are denominator-sensitive", "Field/sub-district yields, crop cuts, compatible product definitions", "High"],
    ["Interannual crop dynamics", "2018-2022 context uses the fixed 2023 crop extent", "No historical crop-area or rotation trend", "Year-specific crop maps and validated annual statistics", "Medium"],
    ["Salinity and economic value", "Not represented in volumetric metrics", "Policy rankings remain incomplete", "Salinity, energy, farm-income, processing, and nutritional indicators", "Medium"],
]
add_table(gaps, ["Gap", "Current evidence", "Consequence", "Evidence needed", "Priority"], gap_rows, 7.5)
gaps.add_heading("Interpretive rule", level=1)
add_p(gaps, "None of these gaps was filled by imputation or unsupported assumptions. The revised manuscript treats them as boundaries on inference and links each one to the analyses or decisions it affects.")
gaps.save(REV / "UNRESOLVED_EVIDENCE_GAPS.docx")


# 5. Supplementary methods and QC.
supp = base_doc("SUPPLEMENTARY METHODS AND QUALITY CONTROL", "Supporting evidence for the revised Indus Basin crop water-accounting manuscript")
supp.add_heading("S1. Analytical scope and reproducibility", level=1)
add_p(supp, "The analytical support is cropland within the Pakistan portion of the Indus drainage basin. It is not an official canal-command delineation. The crop map describes the 2023 seasonal extent: crop-hectares can exceed unique physical area under double cropping, maize combines spring and autumn cycles, and sugarcane is treated as annual. Water accounting is crop-season based and residual blue ET is not measured irrigation withdrawal.")
supp.add_heading("S2. Predictor and classification details", level=1)
add_p(supp, "The Rabi stack contains 42 predictors and the Kharif stack 47. Predictor families include period-specific Sentinel-2 NDVI, EVI, NDWI, LSWI, and SAVI; full-season Sentinel-2 and Landsat optical bands and indices; Sentinel-1 median VV, VH, and VV-VH; and SRTM elevation and slope. Missing predictors were filled with -0.25 before sampling. Both models used 300 trees, minimum leaf population 2, bag fraction 0.70, default variables per split, no maximum-node limit, and classifier seed 20230709. Reference points were assigned by whole 0.05-degree blocks using a deterministic SHA-256 rule with seed 20230709.")
add_csv_table(supp, "Table S1. Class-level validation diagnostics", TABLES / "Table_S1_Classification_Class_Diagnostics.csv")
add_csv_table(supp, "Table S2. Season-level validation diagnostics", TABLES / "Table_S2_Classification_Season_Diagnostics.csv")
supp.add_heading("S3. Crop-area scenario", level=1)
add_p(supp, "The official-area-constrained scenario multiplies mapped mean crop water depth and mapped yield composition by the ratio of official harvested area to mapped crop-season area. It is a first-order structural scenario, not a confidence interval. Because area and production are scaled by the same factor, crop-level VWF remains unchanged while water volume and production scale directly.")
add_csv_table(
    supp,
    "Table S3. Mapped and official-area-constrained crop scenarios",
    TABLES / "Table_S3_Crop_Area_Scenario_Envelope.csv",
    ["crop_name", "mapped_area_ha", "official_harvested_area_ha", "mapped_minus_official_pct", "mapped_total_water_km3", "official_area_scaled_total_water_km3", "mapped_total_vwf_m3_t"],
)
supp.add_heading("S4. Product-chain structural sensitivity", level=1)
add_p(supp, "The primary endpoint combines WaPOR v3 AETI with CHIRPS rainfall; the comparison endpoint combines MODIS ET with ERA5-Land rainfall. Endpoints are reported directly. No uniform interpolation, Monte Carlo distribution, or confidence interval is used.")
add_csv_table(
    supp,
    "Table S4. Product-chain endpoints",
    TABLES / "Table_S4_Product_Chain_Endpoints.csv",
    ["crop_name", "product_chain", "total_water_mm", "green_water_mm", "blue_water_mm", "total_vwf_m3_t", "water_productivity_kg_m3"],
)
add_csv_table(
    supp,
    "Table S5. Direct endpoint structural envelope",
    TABLES / "Table_S5_Product_Chain_Structural_Envelope.csv",
    ["crop_name", "metric", "primary_wapor_chirps", "comparison_modis_era5_land", "absolute_difference", "comparison_minus_primary_pct"],
)
supp.add_heading("S5. Fixed-extent temporal context", level=1)
add_p(supp, "The 2018-2022 baseline and 2023 values use the fixed 2023 crop masks. They describe hydroclimatic/product variation over a short context window, not historical crop-area change or a long-term climatology.")
add_csv_table(supp, "Table S6. Fixed-2023-extent temporal context", TABLES / "Table_S6_2023_Fixed_Extent_Temporal_Context.csv")
supp.add_heading("S6. Hotspot threshold sensitivity", level=1)
add_p(supp, "The district score adds one point each for high blue-water volume, high total VWF, and low WP. Threshold pairs 70/30, 75/25, 80/20, and 90/10 were applied both to all 108 computable districts and to the 76 districts with at least 25,000 mapped crop-hectares. These are internal diagnostic priorities, not environmental or legal thresholds.")
add_csv_table(supp, "Table S7. Hotspot-count sensitivity", TABLES / "Table_S7_Hotspot_Threshold_Sensitivity.csv")
add_csv_table(
    supp,
    "Table S8. District hotspot stability (first 40 rows; complete CSV supplied separately)",
    TABLES / "Table_S8_Hotspot_District_Stability.csv",
    max_rows=40,
)
supp.add_heading("S7. Yield coverage", level=1)
add_p(supp, "Production-dependent indicators were calculated only for positive matched district yields. No unsupported yield was imputed. Of 580 district-crop combinations, 379 were computable and 201 lacked a positive matched yield or mapped crop area.")
add_csv_table(supp, "Table S9. Yield coverage by crop", TABLES / "Table_S9_Yield_Coverage_By_Crop.csv")
add_csv_table(
    supp,
    "Table S10. Yield coverage by district (first 50 rows; complete CSV supplied separately)",
    TABLES / "Table_S10_Yield_Coverage_By_District.csv",
    max_rows=50,
)
supp.add_heading("S8. Numerical QC identities", level=1)
for text in [
    "Green volume + residual-blue volume = total crop water volume within floating-point tolerance.",
    "Mapped crop area x matched district yield = mapped production.",
    "Total water volume / mapped production = total VWF.",
    "Mapped production mass / total water volume = water productivity.",
    "Unavailable yield remains unavailable rather than being displayed as zero.",
]:
    add_p(supp, text, bullet=True)
supp.add_heading("S9. Software and execution record", level=1)
software_rows = [
    ["Operating system", platform.platform()],
    ["Python", sys.version.split()[0]],
    ["Word automation", "Microsoft Word 16.0 COM"],
    ["Geospatial processing", "Google Earth Engine JavaScript/Python APIs; versions recorded in project environment files where available"],
    ["Classification/block seed", "20230709 (authoritative archived output provenance)"],
    ["Structural sensitivity seed", "Not applicable; direct endpoints only"],
    ["Primary CRS", "EPSG:32642; source geographic products retained in their documented CRSs"],
]
add_table(supp, ["Component", "Record"], software_rows, 8.5)
supp.add_heading("S10. Supplementary figures", level=1)
add_p(supp, "Figure S1 is the full 116-district blue-VWF ranking retained outside the main manuscript because journal-scale labels are not readable in a standard figure panel. Figure S2 shows threshold sensitivity for the relative district priority score. High-resolution PNG, TIFF, PDF, and SVG versions are supplied where available.")
supp.save(REV / "SUPPLEMENTARY_METHODS_AND_QC.docx")

print("Created revision reports")
